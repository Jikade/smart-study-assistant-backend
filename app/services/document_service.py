from __future__ import annotations
from sqlalchemy import func, select, text

import pymupdf
import hashlib
import mimetypes
import re
import shutil
import uuid
from datetime import datetime, timezone
from pathlib import Path

from docx import Document as DocxDocument
from fastapi import HTTPException, UploadFile
from pypdf import PdfReader
from sqlalchemy import delete, text
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.db.models import (
    ChunkEmbedding,
    Document,
    DocumentChunk,
    DocumentProcessingJob,
    DocumentSection,
)
from app.services.ai_provider import AIProviderError, get_ai_provider
from app.services.structural_chunker import (
    SSA_SBC_VERSION,
    structural_chunk_text,
)

settings = get_settings()
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


def _safe_name(name: str) -> str:
    name = Path(name).name
    return re.sub(r"[^A-Za-z0-9._-]+", "_", name)[:180] or "document"

def embed_existing_chunks(
    db: Session,
    document_id: int,
    force: bool = False,
) -> int:

    provider = get_ai_provider()

    if not provider.can_embed:
        raise ValueError(
            "Embedding model is not configured. "
            "Check AI_PROVIDER and AI_EMBEDDING_MODEL in .env."
        )

    chunks = list(
        db.scalars(
            select(DocumentChunk)
            .where(
                DocumentChunk.document_id == document_id,
                DocumentChunk.is_active.is_(True),
            )
            .order_by(
                DocumentChunk.chunk_index
            )
        ).all()
    )

    if not chunks:
        raise ValueError(
            "Document has no chunks."
        )

    embeddings_created = 0

    for start in range(0, len(chunks), 32):

        batch = chunks[start:start + 32]

        # Nếu không force thì bỏ qua chunk đã có embedding
        if not force:
            batch = [
                chunk
                for chunk in batch
                if db.scalar(
                    select(ChunkEmbedding.id)
                    .where(
                        ChunkEmbedding.chunk_id
                        == chunk.id
                    )
                )
                is None
            ]

        if not batch:
            continue

        vectors = provider.embeddings(
            [
                chunk.content
                for chunk in batch
            ]
        )

        if len(vectors) != len(batch):
            raise ValueError(
                "Embedding provider returned "
                "unexpected vector count."
            )

        for chunk, vector in zip(
            batch,
            vectors,
        ):

            existing = db.scalar(
                select(ChunkEmbedding)
                .where(
                    ChunkEmbedding.chunk_id
                    == chunk.id
                )
            )

            if existing:
                if not force:
                    continue

                existing.embedding_model = (
                    settings.ai_embedding_model
                    or "unknown"
                )

                existing.embedding_dimension = (
                    len(vector)
                )

                existing.embedding_json = vector

                embedding_row = existing

            else:

                embedding_row = ChunkEmbedding(
                    chunk_id=chunk.id,
                    embedding_model=(
                        settings.ai_embedding_model
                        or "unknown"
                    ),
                    embedding_dimension=len(vector),
                    embedding_json=vector,
                )

                db.add(embedding_row)

            db.flush()

            # Nếu pgvector đã được bật
            if _pgvector_enabled(db):

                vector_literal = (
                    "["
                    + ",".join(
                        f"{float(value):.10g}"
                        for value in vector
                    )
                    + "]"
                )

                db.execute(
                    text(
                        """
                        UPDATE chunk_embeddings
                        SET embedding =
                            CAST(:vector AS vector)
                        WHERE id = :embedding_id
                        """
                    ),
                    {
                        "vector": vector_literal,
                        "embedding_id":
                            embedding_row.id,
                    },
                )

            embeddings_created += 1

    db.commit()

    return embeddings_created


def save_upload(db: Session, owner_id: int, upload: UploadFile, subject_id: int | None) -> Document:
    original = upload.filename or "document"
    ext = Path(original).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(400, f"Unsupported file type: {ext}. Supported: {sorted(ALLOWED_EXTENSIONS)}")

    user_dir = settings.storage_dir / str(owner_id)
    user_dir.mkdir(parents=True, exist_ok=True)
    stored = f"{uuid.uuid4().hex}_{_safe_name(original)}"
    target = user_dir / stored

    sha = hashlib.sha256()
    size = 0
    with target.open("wb") as out:
        while chunk := upload.file.read(1024 * 1024):
            size += len(chunk)
            if size > settings.max_upload_bytes:
                out.close()
                target.unlink(missing_ok=True)
                raise HTTPException(413, f"File exceeds {settings.max_upload_mb} MB")
            sha.update(chunk)
            out.write(chunk)

    mime = upload.content_type or mimetypes.guess_type(original)[0]
    doc = Document(
        owner_id=owner_id,
        subject_id=subject_id,
        original_name=original,
        stored_name=stored,
        storage_url=str(target.resolve()),
        mime_type=mime,
        file_extension=ext.lstrip("."),
        file_size_bytes=size,
        checksum_sha256=sha.hexdigest(),
        status="UPLOADED",
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)
    return doc


def extract_text(path: Path, extension: str) -> tuple[str, int | None]:
    ext = extension.lower().lstrip(".")

    # =========================
    # PDF
    # =========================
    if ext == "pdf":
        doc = pymupdf.open(str(path))

        extracted_pages: list[str] = []

        for page_number, page in enumerate(doc, start=1):

            # Bước 1:
            # Thử lấy text trực tiếp trước
            page_text = page.get_text(
                "text",
                sort=True,
            ).strip()

            # Bước 2:
            # Nếu trang gần như không có text
            # => khả năng cao là scan/image PDF
            if (
                settings.ocr_enabled
                and len(page_text) < settings.ocr_page_min_chars
            ):
                try:
                    text_page = page.get_textpage_ocr(
                        language=settings.ocr_languages,
                        dpi=settings.ocr_dpi,
                        full=True,
                        tessdata=settings.tessdata_dir or None,
                    )

                    page_text = page.get_text(
                        "text",
                        textpage=text_page,
                        sort=True,
                    ).strip()

                except Exception as exc:
                    raise RuntimeError(
                        f"OCR failed at PDF page {page_number}. "
                        f"Check Tesseract installation, PATH, "
                        f"TESSDATA_DIR and language files "
                        f"({settings.ocr_languages}). "
                        f"Original error: {exc}"
                    ) from exc

            if page_text:
                extracted_pages.append(page_text)

        page_count = len(doc)

        doc.close()

        full_text = "\n\n".join(extracted_pages).strip()

        return full_text, page_count

    # =========================
    # DOCX
    # =========================
    if ext == "docx":
        doc = DocxDocument(str(path))

        text_content = "\n".join(
            paragraph.text
            for paragraph in doc.paragraphs
        )

        return text_content, None

    # =========================
    # TXT / Markdown
    # =========================
    if ext in {"txt", "md"}:
        return (
            path.read_text(
                encoding="utf-8",
                errors="ignore",
            ),
            None,
        )

    raise ValueError(
        f"Unsupported extension: {ext}"
    )


def clean_text(value: str) -> str:
    value = value.replace("\x00", " ")
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def chunk_text(value: str, size: int | None = None, overlap: int | None = None) -> list[str]:
    size = size or settings.chunk_size_chars
    overlap = min(overlap if overlap is not None else settings.chunk_overlap_chars, size // 2)
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", value) if p.strip()]
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        candidate = f"{current}\n\n{paragraph}".strip() if current else paragraph
        if len(candidate) <= size:
            current = candidate
            continue
        if current:
            chunks.append(current)
        if len(paragraph) <= size:
            current = paragraph
        else:
            start = 0
            step = max(1, size - overlap)
            while start < len(paragraph):
                part = paragraph[start:start + size].strip()
                if part:
                    chunks.append(part)
                start += step
            current = ""
    if current:
        chunks.append(current)
    return chunks


def _pgvector_enabled(db: Session) -> bool:
    try:
        return bool(db.execute(text("SELECT EXISTS (SELECT 1 FROM pg_extension WHERE extname='vector')")).scalar())
    except Exception:
        return False


def process_document(db: Session, doc: Document) -> tuple[int, int]:
    # Preserve the last usable state before re-indexing.
    previous_status = str(doc.status or "UPLOADED")
    previous_active_count = int(
        db.scalar(
            select(func.count(DocumentChunk.id)).where(
                DocumentChunk.document_id == doc.id,
                DocumentChunk.is_active.is_(True),
            )
        )
        or 0
    )

    job = DocumentProcessingJob(document_id=doc.id, stage="EXTRACT", status="RUNNING", progress_pct=0)
    db.add(job)
    doc.status = "PROCESSING"
    doc.processing_error = None
    db.commit()

    try:
        raw, pages = extract_text(Path(doc.storage_url), doc.file_extension or "")
        cleaned = clean_text(raw)
        if not cleaned:
            raise ValueError("No readable text was extracted from the document")

        job.stage = "CHUNK"
        job.progress_pct = 35

        sections = list(
            db.scalars(
                select(DocumentSection)
                .where(
                    DocumentSection.document_id == doc.id
                )
                .order_by(
                    DocumentSection.section_order,
                    DocumentSection.id,
                )
            ).all()
        )

        structural_chunks = structural_chunk_text(
            cleaned,
            sections=sections,
            size=settings.chunk_size_chars,
            overlap=settings.chunk_overlap_chars,
        )

        if not structural_chunks:
            raise ValueError(
                "SSA-SBC-V1 produced no document chunks"
            )

        chunk_set_id = (
            f"{SSA_SBC_VERSION}-"
            f"{uuid.uuid4().hex[:16]}"
        )

        chunk_rows: list[DocumentChunk] = []

        for index, item in enumerate(structural_chunks):
            row = DocumentChunk(
                document_id=doc.id,
                section_id=item.section_id,
                chunk_index=index,
                chunk_set_id=chunk_set_id,
                chunking_algorithm=SSA_SBC_VERSION,
                is_active=False,
                content=item.content,
                char_count=len(item.content),
                token_count=max(1, len(item.content) // 4),
                content_hash=hashlib.sha256(
                    item.content.encode()
                ).hexdigest(),
                metadata_={
                    "chunking_algorithm": SSA_SBC_VERSION,
                    "chunk_set_id": chunk_set_id,
                    "structural_block_index": item.block_index,
                    "structural_heading": item.structural_heading,
                    "block_char_start": item.block_char_start,
                    "block_char_end": item.block_char_end,
                },
            )
            db.add(row)
            chunk_rows.append(row)

        db.flush()

        embeddings_created = 0
        provider = get_ai_provider()

        if provider.can_embed and chunk_rows:
            job.stage = "EMBED"
            job.progress_pct = 70

            for start in range(0, len(chunk_rows), 32):
                batch = chunk_rows[start:start + 32]
                vectors = provider.embeddings(
                    [chunk.content for chunk in batch]
                )

                if len(vectors) != len(batch):
                    raise ValueError(
                        "Embedding provider returned "
                        "an unexpected number of vectors"
                    )

                for chunk, vector in zip(batch, vectors):
                    emb = ChunkEmbedding(
                        chunk_id=chunk.id,
                        embedding_model=(
                            settings.ai_embedding_model
                            or "unknown"
                        ),
                        embedding_dimension=len(vector),
                        embedding_json=vector,
                    )
                    db.add(emb)
                    db.flush()

                    if _pgvector_enabled(db):
                        vector_literal = (
                            "["
                            + ",".join(
                                f"{float(value):.10g}"
                                for value in vector
                            )
                            + "]"
                        )
                        db.execute(
                            text(
                                "UPDATE chunk_embeddings "
                                "SET embedding = CAST(:v AS vector) "
                                "WHERE id = :id"
                            ),
                            {
                                "v": vector_literal,
                                "id": emb.id,
                            },
                        )

                    embeddings_created += 1

        if (
            provider.can_embed
            and embeddings_created != len(chunk_rows)
        ):
            raise ValueError(
                "New chunk set is incomplete: "
                "not every chunk has an embedding"
            )

        hard_heading_re = re.compile(
            r"(?im)^\s*"
            r"(?:CHƯƠNG|CHUONG|CHAPTER|PHẦN|PHAN|PART)"
            r"\s+(?:\d+|[IVXLCDM]+)"
        )

        for row in chunk_rows:
            if len(
                hard_heading_re.findall(row.content)
            ) > 1:
                raise ValueError(
                    "SSA-SBC-V1 boundary validation failed: "
                    "one chunk contains multiple hard "
                    "structural headings"
                )

        switched_at = datetime.now(timezone.utc)

        db.query(DocumentChunk).filter(
            DocumentChunk.document_id == doc.id,
            DocumentChunk.is_active.is_(True),
        ).update(
            {
                DocumentChunk.is_active: False,
                DocumentChunk.superseded_at: switched_at,
            },
            synchronize_session=False,
        )

        db.query(DocumentChunk).filter(
            DocumentChunk.document_id == doc.id,
            DocumentChunk.chunk_set_id == chunk_set_id,
        ).update(
            {
                DocumentChunk.is_active: True,
                DocumentChunk.superseded_at: None,
            },
            synchronize_session=False,
        )

        db.flush()

        active_count = db.scalar(
            select(func.count(DocumentChunk.id)).where(
                DocumentChunk.document_id == doc.id,
                DocumentChunk.chunk_set_id == chunk_set_id,
                DocumentChunk.is_active.is_(True),
            )
        )

        if int(active_count or 0) != len(chunk_rows):
            raise ValueError(
                "Atomic chunk-set switch validation failed"
            )

        doc.page_count = pages
        doc.status = "READY"
        doc.processed_at = datetime.now(timezone.utc)
        job.stage = "COMPLETE"
        job.status = "SUCCEEDED"
        job.progress_pct = 100
        job.completed_at = datetime.now(timezone.utc)
        db.commit()
        return len(chunk_rows), embeddings_created
    except Exception as exc:
        db.rollback()
        doc = db.get(Document, doc.id)
        if doc:
            active_count_after_rollback = int(
                db.scalar(
                    select(func.count(DocumentChunk.id)).where(
                        DocumentChunk.document_id == doc.id,
                        DocumentChunk.is_active.is_(True),
                    )
                )
                or 0
            )

            if (
                previous_status == "READY"
                and previous_active_count > 0
                and active_count_after_rollback > 0
            ):
                doc.status = "READY"
            else:
                doc.status = "FAILED"

            doc.processing_error = str(exc)
        job = db.get(DocumentProcessingJob, job.id)
        if job:
            job.status = "FAILED"
            job.error_message = str(exc)
            job.completed_at = datetime.now(timezone.utc)
        db.commit()
        raise
